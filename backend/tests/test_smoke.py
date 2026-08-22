"""不依赖 MinerU 引擎与模型的冒烟测试：用假转换器跑通完整流程。"""
from __future__ import annotations

import io
import threading
import time
import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.app import api, config
from backend.app.converter import ConversionError, ConversionResult, Converter, MinerUConverter
from backend.app.jobs import JobManager
from backend.app.main import create_app

PDF_BYTES = b"%PDF-1.4\n% fake but magic-valid\n"


class FakeConverter(Converter):
    """模拟 MinerU 的输出结构：<out>/<stem>/<method>/<stem>.md + images/。"""

    def convert(self, pdf_path, output_dir, *, lang="ch", force_ocr=False):
        method = "ocr" if force_ocr else "auto"
        md_dir = output_dir / pdf_path.stem / method
        images = md_dir / "images"
        images.mkdir(parents=True, exist_ok=True)
        (images / "pic.png").write_bytes(b"\x89PNG\r\n\x1a\n")
        md = md_dir / f"{pdf_path.stem}.md"
        md.write_text(
            "# 标题\n\n![pic](images/pic.png)\n\n公式 $E=mc^2$\n", encoding="utf-8"
        )
        return ConversionResult(markdown_path=md, images_dir=images, output_dir=output_dir)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    monkeypatch.setattr(config, "OUTPUT_DIR", tmp_path / "outputs")
    config.ensure_dirs()
    app = create_app()
    api.init(JobManager(FakeConverter()))
    with TestClient(app) as c:
        yield c


def _convert(client: TestClient, name="doc.pdf", data=PDF_BYTES, **form):
    files = {"file": (name, data, "application/pdf")}
    return client.post("/api/convert", files=files, data=form)


def _wait_finished(client: TestClient, job_id: str, timeout: float = 10.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        s = client.get(f"/api/status/{job_id}").json()
        if s["status"] in {"done", "error"}:
            return s
        time.sleep(0.05)
    raise AssertionError("job did not finish in time")


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_reject_non_pdf_extension(client):
    r = _convert(client, name="a.txt", data=b"%PDF-1.4")
    assert r.status_code == 400


def test_reject_bad_magic(client):
    r = _convert(client, data=b"not a pdf at all")
    assert r.status_code == 400


def test_reject_oversize(client, monkeypatch):
    monkeypatch.setattr(config, "MAX_UPLOAD_BYTES", 4)
    r = _convert(client)
    assert r.status_code == 413


def test_status_unknown_job_404(client):
    assert client.get("/api/status/nope").status_code == 404


def test_full_conversion_flow(client):
    r = _convert(client, lang="ch", force_ocr="false")
    assert r.status_code == 200
    job_id = r.json()["job_id"]

    s = _wait_finished(client, job_id)
    assert s["status"] == "done", s
    assert s["progress"] == 100

    res = client.get(f"/api/result/{job_id}")
    assert res.status_code == 200
    body = res.json()
    assert body["filename"] == "doc.pdf"
    assert "# 标题" in body["markdown"]
    # 图片链接已改写到 /api/files
    assert f"](/api/files/{job_id}/images/pic.png)" in body["markdown"]
    assert body["images"] == ["pic.png"]

    img = client.get(f"/api/files/{job_id}/images/pic.png")
    assert img.status_code == 200
    assert img.content.startswith(b"\x89PNG")

    src = client.get(f"/api/source/{job_id}")
    assert src.status_code == 200
    assert src.headers["content-type"] == "application/pdf"

    dl = client.get(f"/api/download/{job_id}?format=md")
    assert dl.status_code == 200
    assert b"# \xe6\xa0\x87\xe9\xa2\x98" in dl.content  # "# 标题" 的 UTF-8

    z = client.get(f"/api/download/{job_id}?format=zip")
    assert z.status_code == 200
    zf = zipfile.ZipFile(io.BytesIO(z.content))
    names = zf.namelist()
    assert any(n.endswith(".md") for n in names)
    assert "images/pic.png" in names


def test_result_before_done_409(client):
    event = threading.Event()

    class BlockingConverter(FakeConverter):
        def convert(self, *args, **kwargs):
            event.wait(5)
            return super().convert(*args, **kwargs)

    api.init(JobManager(BlockingConverter()))
    job_id = _convert(client).json()["job_id"]
    try:
        assert client.get(f"/api/result/{job_id}").status_code == 409
        assert client.get(f"/api/download/{job_id}").status_code == 409
    finally:
        event.set()
    assert _wait_finished(client, job_id)["status"] == "done"


def test_engine_error_surfaces_in_status(client):
    class FailingConverter(FakeConverter):
        def convert(self, *args, **kwargs):
            raise ConversionError("boom: engine exploded")

    api.init(JobManager(FailingConverter()))
    job_id = _convert(client).json()["job_id"]
    s = _wait_finished(client, job_id)
    assert s["status"] == "error"
    assert "boom" in s["message"]


def test_path_traversal_blocked(client):
    job_id = _convert(client).json()["job_id"]
    _wait_finished(client, job_id)
    r = client.get(f"/api/files/{job_id}/..%2F..%2Fuploads%2F{job_id}.pdf")
    assert r.status_code == 404


def test_find_markdown_prefers_same_stem(tmp_path):
    other = tmp_path / "x" / "auto"
    other.mkdir(parents=True)
    (other / "other.md").write_text("big" * 1000, encoding="utf-8")
    same = tmp_path / "doc" / "auto"
    same.mkdir(parents=True)
    target = same / "doc.md"
    target.write_text("small", encoding="utf-8")
    assert MinerUConverter._find_markdown(tmp_path, "doc") == target


def test_find_markdown_falls_back_to_largest(tmp_path):
    d = tmp_path / "x" / "auto"
    d.mkdir(parents=True)
    (d / "a.md").write_text("tiny", encoding="utf-8")
    big = d / "b.md"
    big.write_text("big" * 1000, encoding="utf-8")
    assert MinerUConverter._find_markdown(tmp_path, "doc") == big


def test_find_markdown_none(tmp_path):
    assert MinerUConverter._find_markdown(tmp_path, "doc") is None


def test_executable_env_override(tmp_path, monkeypatch):
    fake_exe = tmp_path / "mineru.exe"
    fake_exe.write_text("x", encoding="utf-8")
    monkeypatch.setenv("ANY2MD_MINERU_EXE", str(fake_exe))
    assert MinerUConverter()._executable() == str(fake_exe)


def test_executable_missing_raises(monkeypatch, tmp_path):
    monkeypatch.delenv("ANY2MD_MINERU_EXE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _: None)
    monkeypatch.setattr("sys.executable", str(tmp_path / "python.exe"))
    with pytest.raises(ConversionError, match="not installed"):
        MinerUConverter()._executable()


# ---------------- Word (.docx) 支持 ----------------

def _make_docx(text: str) -> bytes:
    """现场生成一个最小可用的 .docx（zip 容器）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        zf.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        zf.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>",
        )
    return buf.getvalue()


def test_get_converter_for_routing():
    from backend.app.converter import WordConverter, get_converter_for

    assert isinstance(get_converter_for("a.docx"), WordConverter)
    assert isinstance(get_converter_for("a.doc"), WordConverter)
    assert isinstance(get_converter_for("a.pdf"), MinerUConverter)


def test_word_converter_real(tmp_path):
    """用真实 mammoth 转换一个最小 docx，验证产出 Markdown。"""
    from backend.app.converter import WordConverter

    pytest.importorskip("mammoth")
    docx = tmp_path / "hello.docx"
    docx.write_bytes(_make_docx("你好 Word"))
    out = tmp_path / "out"
    res = WordConverter().convert(docx, out)
    assert res.markdown_path.is_file()
    text = res.markdown_path.read_text(encoding="utf-8")
    assert "你好 Word" in text


def test_convert_docx_via_api(client):
    """docx 走完整 API 流程（client fixture 用 FakeConverter，绕过真实引擎）。"""
    r = _convert(client, name="doc.docx", data=_make_docx("hi"))
    assert r.status_code == 200
    job_id = r.json()["job_id"]
    s = _wait_finished(client, job_id)
    assert s["status"] == "done", s
    # 源文件按 docx media type 返回
    src = client.get(f"/api/source/{job_id}")
    assert src.status_code == 200
    assert "wordprocessingml" in src.headers["content-type"]


def test_reject_unsupported_extension(client):
    r = _convert(client, name="a.txt", data=b"hello")
    assert r.status_code == 400


def test_reject_bad_docx_magic(client):
    r = _convert(client, name="a.docx", data=b"not a zip")
    assert r.status_code == 400


# ---------------- PowerPoint (.pptx) 支持 ----------------

def _make_pptx() -> bytes:
    """现场生成一个最小可用的 .pptx：标题 + 副标题 + 引导句 + 两级项目符号 + 表格。"""
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # 标题和内容
    slide.shapes.title.text = "季度汇报"
    body = slide.placeholders[1].text_frame
    body.text = "营收增长"
    p = body.add_paragraph()
    p.text = "同比增长 20%"
    p.level = 1
    # 副标题幻灯片（标题 + 副标题布局）
    s2 = prs.slides.add_slide(prs.slide_layouts[0])  # 标题幻灯片：标题 + 副标题
    s2.shapes.title.text = "仿真挂死的场景全景"
    s2.placeholders[1].text_frame.text = "挂死 = 仿真时间不再推进；先分类，再对症下药"
    # 表格
    gfx = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    tbl = gfx.table
    tbl.cell(0, 0).text = "指标"
    tbl.cell(0, 1).text = "数值"
    tbl.cell(1, 0).text = "营收"
    tbl.cell(1, 1).text = "100"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_get_converter_for_routing_pptx():
    from backend.app.converter import MinerUConverter, PptxConverter, get_converter_for

    assert isinstance(get_converter_for("a.pptx"), PptxConverter)
    assert isinstance(get_converter_for("a.ppt"), PptxConverter)
    assert isinstance(get_converter_for("a.pdf"), MinerUConverter)


def test_pptx_converter_real(tmp_path):
    """用真实 python-pptx 转换一个最小 pptx，验证标题/副标题/列表/表格保持不变。"""
    from backend.app.converter import PptxConverter

    pytest.importorskip("pptx")
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(_make_pptx())
    res = PptxConverter().convert(pptx_path, tmp_path / "out")
    assert res.markdown_path.is_file()
    text = res.markdown_path.read_text(encoding="utf-8")
    assert "## 季度汇报" in text  # 标题 → 二级标题
    assert "## 仿真挂死的场景全景" in text
    assert "### 挂死 = 仿真时间不再推进；先分类，再对症下药" in text  # 副标题 → 三级标题
    assert "- 营收增长" in text  # 一级项目符号
    assert "  1. 同比增长 20%" in text or "  - 同比增长 20%" in text  # 二级缩进
    assert "| 指标 | 数值 |" in text  # 表格表头
    assert "| 营收 | 100 |" in text  # 表格内容


def test_pptx_legacy_ppt_clear_error(tmp_path):
    """老式 .ppt（OLE2）给出明确报错，提示另存为 .pptx。"""
    from backend.app.converter import ConversionError, PptxConverter

    with pytest.raises(ConversionError, match="save as .pptx"):
        PptxConverter().convert(tmp_path / "old.ppt", tmp_path / "out")


def test_convert_pptx_via_api(client):
    """pptx 走完整 API 流程（client fixture 用 FakeConverter，绕过真实引擎）。"""
    r = _convert(client, name="deck.pptx", data=_make_pptx())
    assert r.status_code == 200
    job_id = r.json()["job_id"]
    s = _wait_finished(client, job_id)
    assert s["status"] == "done", s
    # 源文件按 pptx media type 返回
    src = client.get(f"/api/source/{job_id}")
    assert src.status_code == 200
    assert "presentationml" in src.headers["content-type"]


def test_reject_bad_pptx_magic(client):
    r = _convert(client, name="a.pptx", data=b"not a zip")
    assert r.status_code == 400
